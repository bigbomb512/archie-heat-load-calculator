"""Read-only, content-addressed equipment catalogue. No project or network I/O.

Two workbooks: data/equipment_catalogue.xlsx (442 sourced examples of specific models and states, used when the
drawings print a matching model) and data/equipment_typical_values.xlsx (typical heat per equipment type and size,
with the use pattern of each, used when they don't). Both are edited by people and only read here."""

from collections import Counter
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import hashlib
import math
import re
import threading

DEFAULT_PATH = Path(__file__).resolve().parents[1] / "data/equipment_catalogue.xlsx"
TYPICAL_PATH = Path(__file__).resolve().parents[1] / "data/equipment_typical_values.xlsx"
RESEARCH_NOTE = "Further research required if this item is used in the project."
HEADERS = {
    "id": "Equipment ID", "name": "Object", "aliases": "Matching aliases (semicolon separated)",
    "eligibility": "Calculation eligibility", "sensible_w": "Conditional sensible heat (W/unit)",
    "latent_w": "Conditional latent heat (W/unit)",
    "conditions": "Conditions / missing information — read before using", "source": "Source",
    "model": "Manufacturer / model", "state": "Operating state / basis",
    "heat_location": "Heat rejection location", "category": "Category",
    "confidence": "Confidence / evidence",
}
ELIGIBILITY = {"CONDITIONAL_MATCH", "ESTIMATE_ONLY", "REFERENCE_ONLY"}
_CACHE = {}
_LOCK = threading.Lock()


def normalise(value):
    text = re.sub(r"\bu/b\b|\bub\b|\bunder[ -]bench\b", "underbench", str(value or "").casefold())
    text = re.sub(r"\bfridges?\b", "refrigerator", text)
    text = re.sub(r"\bice makers?\b", "ice machine", text)
    text = re.sub(r"\bdisp\b", "display", text)
    text = re.sub(r"\bsign\b", "signage", text)
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


def number(value):
    return (float(value) if isinstance(value, (int, float)) and not isinstance(value, bool)
            and math.isfinite(value) and value >= 0 else None)


def load(path=None):
    path = Path(path or DEFAULT_PATH)
    try:
        raw = path.read_bytes()
    except OSError as error:
        return {"hash": "unavailable", "records": [], "diagnostics": [],
                "error": f"Equipment catalogue cannot be read: {error.strerror or type(error).__name__}."}
    digest = hashlib.sha256(raw).hexdigest()
    with _LOCK:
        if digest in _CACHE:
            return deepcopy(_CACHE[digest])
    result = {"hash": digest, "records": [], "diagnostics": [], "error": ""}
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False)
        try:
            sheet = workbook["Approximate heat gains"]
            iterator = sheet.iter_rows(values_only=True)
            for row_number, row in enumerate(iterator, 1):
                columns = {str(value).strip(): i for i, value in enumerate(row) if value is not None}
                if HEADERS["id"] in columns:
                    break
            else:
                raise ValueError("Equipment header row not found.")
            missing = set(HEADERS.values()) - columns.keys()
            if missing:
                raise ValueError("Missing catalogue columns: " + ", ".join(sorted(missing)))
            candidates = []
            for row_number, row in enumerate(iterator, row_number + 1):
                record = {key: row[columns[header]] if columns[header] < len(row) else None
                          for key, header in HEADERS.items()}
                if not any(value is not None and value != "" for value in record.values()):
                    continue
                record.update({"row": row_number, "sheet": sheet.title})
                candidates.append(record)
            counts = Counter(str(row["id"] or "").strip() for row in candidates)
            for row in candidates:
                errors = []
                for field in ("id", "name", "conditions", "source", "eligibility"):
                    if not isinstance(row[field], str) or not row[field].strip():
                        errors.append(f"Missing {field}.")
                row["id"] = str(row["id"] or "").strip()
                if counts[row["id"]] > 1:
                    errors.append("Duplicate equipment ID.")
                if row["eligibility"] not in ELIGIBILITY:
                    errors.append("Unsupported calculation eligibility.")
                if not str(row["source"] or "").startswith("https://"):
                    errors.append("Source must be an HTTPS reference.")
                for field in ("sensible_w", "latent_w"):
                    if row[field] == "":
                        row[field] = None
                    if row[field] is not None and number(row[field]) is None:
                        errors.append(f"{field} must be finite, nonnegative numeric watts per unit.")
                    if row["eligibility"] == "CONDITIONAL_MATCH" and number(row[field]) is None:
                        errors.append(f"Conditional entry missing {field}.")
                if errors:
                    result["diagnostics"].append({"id": row["id"], "row": row["row"], "reasons": errors})
                    continue
                row["aliases"] = sorted({normalise(row["name"]), *[normalise(x) for x in str(row["aliases"] or "").split(";")]} - {""})
                result["records"].append(row)
        finally:
            workbook.close()
    except Exception as error:
        result.update(records=[], error=f"Equipment catalogue cannot be loaded: {type(error).__name__}: {error}")
    with _LOCK:
        if len(_CACHE) >= 4:
            _CACHE.pop(next(iter(_CACHE)))
        _CACHE[digest] = deepcopy(result)
    return result


def candidates(catalogue, item):
    """Candidate discovery only; even an exact name does not prove applicability."""
    names = {normalise(item.get(key)) for key in ("name", "model", "manufacturer_model")} - {""}
    exact = [row for row in catalogue["records"] if names.intersection(row["aliases"])
             or normalise(row.get("model")) in names]
    if exact:
        return exact
    tokens = {word for word in normalise(item.get("name")).split() if not word.isdigit()} - {
        "small", "medium", "large", "commercial", "electric", "single", "door", "drink", "tier", "inch", "new", "wall", "mounted", "illuminated"}
    ranked = []
    for row in catalogue["records"]:
        other = set(normalise(row["name"]).split())
        score = max(len(tokens & other) / max(1, len(tokens | other)), .5 * len(tokens & other) / max(1, len(tokens)))
        if score >= .35:
            ranked.append((score, row))
    return [row for _, row in sorted(ranked, key=lambda pair: (-pair[0], pair[1]["id"]))[:5]]


TYPICAL_HEADERS = {
    "id": "Typical ID", "type": "Equipment type", "band": "Size band",
    "size_rule": "How to tell the size from the drawings", "aliases": "Matching aliases (semicolon separated)",
    "generic_for": "Generic fallback for type",
    "sensible_w": "Sensible W per unit (free-standing)", "latent_w": "Latent W per unit (free-standing)",
    "hood_sensible_w": "Sensible W per unit (under a hood)", "hood_latent_w": "Latent W per unit (under a hood)",
    "pattern": "Use pattern", "heat_location": "Heat rejection location", "based_on": "Based on (catalogue IDs)",
    "source": "Source",
}
PATTERN_HEADERS = {"id": "Pattern ID", "description": "Description", "basis": "Basis",
                   "fraction": "Fraction of the per-unit heat (blank = the room type's equipment diversity)"}
PATTERN_BASES = {"CONTINUOUS", "ROOM_HOURS"}
TYPES = {"cooking", "refrigeration", "dishwashing", "hot_drinks", "electronics", "signage", "other"}


def _rows(workbook, sheet_name, headers):
    """Rows of a sheet as dicts by header name, found by the first header; (row number, dict) pairs."""
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    first = next(iter(headers.values()))
    for row_number, row in enumerate(iterator, 1):
        columns = {str(value).strip(): i for i, value in enumerate(row) if value is not None}
        if first in columns:
            break
    else:
        raise ValueError(f"{sheet_name}: header row not found.")
    missing = set(headers.values()) - columns.keys()
    if missing:
        raise ValueError(f"{sheet_name}: missing columns: " + ", ".join(sorted(missing)))
    out = []
    for row_number, row in enumerate(iterator, row_number + 1):
        record = {key: row[columns[header]] if columns[header] < len(row) else None for key, header in headers.items()}
        if any(value is not None and value != "" for value in record.values()):
            out.append((row_number, record))
    return out


def load_typical(path=None):
    """Typical values and use patterns. A row whose free-standing watts are blank is not filled in yet: it is listed
    in `waiting`, not used and not an error. Invalid filled rows are quarantined with their reasons."""
    path = Path(path or TYPICAL_PATH)
    try:
        raw = path.read_bytes()
    except OSError as error:
        return {"hash": "unavailable", "records": [], "patterns": {}, "waiting": [], "diagnostics": [],
                "error": f"Typical values workbook cannot be read: {error.strerror or type(error).__name__}."}
    digest = hashlib.sha256(raw).hexdigest()
    key = ("typical", digest)
    with _LOCK:
        if key in _CACHE:
            return deepcopy(_CACHE[key])
    result = {"hash": digest, "records": [], "patterns": {}, "waiting": [], "diagnostics": [], "error": ""}
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False)
        try:
            for row_number, row in _rows(workbook, "Use patterns", PATTERN_HEADERS):
                pattern_id = str(row["id"] or "").strip()
                fraction = row["fraction"] if row["fraction"] not in ("", None) else None
                if row["basis"] not in PATTERN_BASES or (fraction is not None and (number(fraction) is None or fraction > 1)):
                    result["diagnostics"].append({"id": pattern_id, "row": row_number, "sheet": "Use patterns",
                                                  "reasons": ["Basis must be CONTINUOUS or ROOM_HOURS and the fraction 0–1 or blank."]})
                    continue
                result["patterns"][pattern_id] = {"id": pattern_id, "basis": row["basis"], "fraction": number(fraction),
                                                  "description": str(row["description"] or "")}
            candidates = _rows(workbook, "Typical values", TYPICAL_HEADERS)
        finally:
            workbook.close()
        counts = Counter(str(row["id"] or "").strip() for _, row in candidates)
        for row_number, row in candidates:
            row = {key: (None if value == "" else value) for key, value in row.items()}
            row.update(id=str(row["id"] or "").strip(), row=row_number, sheet="Typical values")
            if row["sensible_w"] is None and row["latent_w"] is None:
                result["waiting"].append({"id": row["id"], "type": row["type"], "band": row["band"]})
                continue
            errors = []
            if not row["id"] or counts[row["id"]] > 1:
                errors.append("Missing or duplicate Typical ID.")
            if not isinstance(row["type"], str) or not row["type"].strip():
                errors.append("Missing equipment type.")
            for field in ("sensible_w", "latent_w"):
                if number(row[field]) is None:
                    errors.append(f"{TYPICAL_HEADERS[field]} must be a number of watts (0 or more).")
            for field in ("hood_sensible_w", "hood_latent_w"):
                if row[field] is not None and number(row[field]) is None:
                    errors.append(f"{TYPICAL_HEADERS[field]} must be blank or a number of watts.")
            if (row["hood_sensible_w"] is None) != (row["hood_latent_w"] is None):
                errors.append("Fill both under-hood values or neither.")
            if row["pattern"] not in result["patterns"]:
                errors.append("Use pattern must be one of the Use patterns sheet's IDs.")
            if row["generic_for"] is not None and row["generic_for"] not in TYPES:
                errors.append("Generic fallback for type must be one of: " + ", ".join(sorted(TYPES)) + ".")
            if not str(row["source"] or "").startswith("https://"):
                errors.append("Source must be an HTTPS reference.")
            if errors:
                result["diagnostics"].append({"id": row["id"], "row": row_number, "sheet": "Typical values", "reasons": errors})
                continue
            row["name"] = f"{row['type'].strip()} ({row['band']})" if row["band"] else row["type"].strip()
            row["aliases"] = sorted({normalise(row["type"]), *[normalise(x) for x in str(row["aliases"] or "").split(";")]} - {""})
            result["records"].append(row)
    except Exception as error:
        result.update(records=[], patterns={}, error=f"Typical values workbook cannot be loaded: {type(error).__name__}: {error}")
    with _LOCK:
        if len(_CACHE) >= 8:
            _CACHE.pop(next(iter(_CACHE)))
        _CACHE[key] = deepcopy(result)
    return result


def current_hash(book=None, typical=None):
    """One hash for both workbooks: a change to either makes a saved Phase 2 result out of date."""
    book = book or load()
    typical = typical or load_typical()
    return hashlib.sha256(f"{book['hash']}:{typical['hash']}".encode()).hexdigest()


def typical_candidates(typical, item):
    """Filled typical rows whose type or aliases match the item's name (every size band of that type)."""
    name = normalise(item.get("name"))
    words = set(name.split())
    found = []
    for row in typical["records"]:
        if row.get("generic_for"):
            continue
        if any(alias == name or (alias and f" {alias} " in f" {name} ") or (len(alias.split()) > 1 and set(alias.split()) <= words)
               for alias in row["aliases"]):
            found.append(row)
    return found


def generic_row(typical, kind):
    return next((row for row in typical["records"] if row.get("generic_for") == kind), None)
