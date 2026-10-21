"""Equipment catalogue and typical values: loading, the value ladder (data sheet, matched model, typical, generic),
the hood choice, never stopping, and the room-load integration. No provider calls."""
from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from ai import equipment_catalogue as cat, ai_preliminary, room_use_resolution, internal_gains_resolution
from ai.heat_loads import equipment_load
from backend import equipment_catalogue_service as service, ai_preliminary_service, page_extraction_service


def schedule(fraction=.5):
    return {day: [fraction] * 24 for day in ("weekday", "saturday", "sunday", "holiday")}


def record(**changes):
    return {"id": "EQ-1", "name": "Test PC", "aliases": "test desktop; PC model A", "eligibility": "CONDITIONAL_MATCH",
            "sensible_w": 80, "latent_w": 20, "conditions": "Model A running locally, no exported heat. Hourly state basis.",
            "source": "https://manufacturer.example/model-a", "model": "PC model A", "state": "active",
            "heat_location": "room", "category": "Computers", **changes}


def workbook(path, rows=None, reverse=False):
    w = Workbook();s = w.active;s.title = "Approximate heat gains"
    keys = list(cat.HEADERS)
    if reverse:
        keys.reverse()
    s.append([cat.HEADERS[k] for k in keys])
    for row in rows or [record()]:
        s.append([row.get(k) for k in keys])
    w.save(path);w.close()


TYPICAL_ROWS = [
    {"id": "TV-1", "type": "Under-bench fridge", "band": "1-door", "aliases": "UB fridge;UBR", "sensible_w": 350, "latent_w": 0, "pattern": "CONTINUOUS"},
    {"id": "TV-2", "type": "Under-bench fridge", "band": "2-door", "aliases": "UB fridge;UBR", "sensible_w": 500, "latent_w": 0, "pattern": "CONTINUOUS"},
    {"id": "TV-3", "type": "Deep fryer", "band": "Countertop", "aliases": "fryer", "sensible_w": 700, "latent_w": 1600,
     "hood_sensible_w": 300, "hood_latent_w": 0, "pattern": "WHILE_OPEN"},
    {"id": "TV-9", "type": "Any cooking appliance", "band": "Generic", "generic_for": "cooking", "sensible_w": 1000, "latent_w": 500, "pattern": "WHILE_OPEN"},
]
PATTERNS = [{"id": "CONTINUOUS", "description": "On around the clock", "basis": "CONTINUOUS", "fraction": 1.0},
            {"id": "WHILE_OPEN", "description": "While open", "basis": "ROOM_HOURS", "fraction": None}]


def typical_workbook(path, rows=TYPICAL_ROWS, patterns=PATTERNS):
    w = Workbook(); s = w.active; s.title = "Typical values"
    s.append(["Instructions above the header are ignored."]); s.append(list(cat.TYPICAL_HEADERS.values()))
    for row in rows:
        s.append([{"source": "https://example.org/typical"}.get(key) if key not in row else row[key] for key in cat.TYPICAL_HEADERS])
    p = w.create_sheet("Use patterns"); p.append(list(cat.PATTERN_HEADERS.values()))
    for row in patterns:
        p.append([row.get(key) for key in cat.PATTERN_HEADERS])
    w.save(path); w.close()


class CatalogueTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "catalogue.xlsx"

    def test_shipped_workbook_and_ids(self):
        value = cat.load()
        self.assertFalse(value["error"])
        self.assertEqual(len(value["records"]), 442)
        self.assertEqual(value["diagnostics"], [])

    def test_reordered_headers_added_rows_and_aliases(self):
        workbook(self.path, [record(), record(id="EQ-2", name="Other")], reverse=True)
        result = cat.load(self.path)
        self.assertEqual(len(result["records"]), 2)
        self.assertEqual(cat.candidates(result, {"name": "TEST DESKTOP"})[0]["id"], "EQ-1")

    def test_duplicate_ids_quarantine_all_colliding_rows(self):
        workbook(self.path, [record(), record(name="Duplicate"), record(id="EQ-3")])
        result = cat.load(self.path)
        self.assertEqual([r["id"] for r in result["records"]], ["EQ-3"])
        self.assertEqual(len(result["diagnostics"]), 2)

    def test_invalid_rows_do_not_hide_valid_rows(self):
        workbook(self.path, [record(), record(id="bad", sensible_w=-1), record(id="bad2", source=""),
                              record(id="bad3", eligibility="AUTOMATIC")])
        result = cat.load(self.path)
        self.assertEqual(len(result["records"]), 1)
        self.assertEqual(len(result["diagnostics"]), 3)

    def test_blank_is_unknown_but_zero_is_numeric(self):
        workbook(self.path, [record(sensible_w=0, latent_w=40), record(id="blank", sensible_w=None),
                              record(id="reference", sensible_w="", latent_w="", eligibility="REFERENCE_ONLY")])
        result = cat.load(self.path)
        self.assertEqual([r["id"] for r in result["records"]], ["EQ-1", "reference"])
        self.assertIsNone(result["records"][1]["sensible_w"])

    def test_hash_changes_and_missing_or_corrupt_file_is_explicit(self):
        workbook(self.path);first = cat.load(self.path)
        workbook(self.path, [record(sensible_w=200)])
        self.assertNotEqual(first["hash"], cat.load(self.path)["hash"])
        self.path.unlink();self.assertTrue(cat.load(self.path)["error"])
        self.path.write_bytes(b"not an xlsx");self.assertTrue(cat.load(self.path)["error"])

    def test_required_header_and_formula_heat_rejected(self):
        workbook(self.path, [record(sensible_w="=100/2")])
        self.assertEqual(len(cat.load(self.path)["diagnostics"]), 1)
        w = Workbook();w.save(self.path);w.close()
        self.assertTrue(cat.load(self.path)["error"])


class TypicalValuesTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "typical.xlsx"

    def test_shipped_sheet_lists_the_types_with_values_still_to_fill(self):
        value = cat.load_typical()
        self.assertFalse(value["error"])
        self.assertEqual(set(value["patterns"]), {"CONTINUOUS", "WHILE_OPEN"})
        self.assertGreaterEqual(len(value["records"]) + len(value["waiting"]), 60)
        self.assertEqual(value["diagnostics"], [])

    def test_filled_rows_load_and_blank_ones_wait(self):
        typical_workbook(self.path, TYPICAL_ROWS + [{"id": "TV-5", "type": "Chest freezer", "band": "Any", "pattern": "CONTINUOUS"}])
        value = cat.load_typical(self.path)
        self.assertEqual([row["id"] for row in value["records"]], ["TV-1", "TV-2", "TV-3", "TV-9"])
        self.assertEqual(value["waiting"], [{"id": "TV-5", "type": "Chest freezer", "band": "Any"}])
        self.assertEqual([row["id"] for row in cat.typical_candidates(value, {"name": "UB FRIDGE"})], ["TV-1", "TV-2"])
        self.assertEqual(cat.generic_row(value, "cooking")["id"], "TV-9")

    def test_invalid_rows_are_quarantined_with_reasons(self):
        bad = [{**TYPICAL_ROWS[0], "id": "TV-A", "sensible_w": -1}, {**TYPICAL_ROWS[0], "id": "TV-B", "pattern": "SOMETIMES"},
               {**TYPICAL_ROWS[2], "id": "TV-C", "hood_latent_w": None}, {**TYPICAL_ROWS[0], "id": "TV-D", "source": "www.x"},
               {**TYPICAL_ROWS[0], "id": "TV-E", "generic_for": "boats"}, {**TYPICAL_ROWS[0], "id": "TV-1"}]
        value = cat.load_typical(self.path) if typical_workbook(self.path, TYPICAL_ROWS + bad) is None else None
        self.assertEqual([row["id"] for row in value["records"]], ["TV-2", "TV-3", "TV-9"])   # both TV-1s are out
        self.assertEqual(len(value["diagnostics"]), 7)

    def test_either_workbook_changing_changes_the_hash_and_a_missing_one_is_explicit(self):
        typical_workbook(self.path)
        first = cat.current_hash(cat.load(), cat.load_typical(self.path))
        typical_workbook(self.path, [{**TYPICAL_ROWS[0], "sensible_w": 360}])
        self.assertNotEqual(first, cat.current_hash(cat.load(), cat.load_typical(self.path)))
        self.path.unlink()
        self.assertIn("cannot be read", cat.load_typical(self.path)["error"])


def check(quote, page=2):
    return {"satisfied": True, "reason": "Shown on the drawing.", "evidence": [{"page": page, "quote": quote}]}


class LadderTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        workbook(self.root / "catalogue.xlsx"); self.book = cat.load(self.root / "catalogue.xlsx")
        typical_workbook(self.root / "typical.xlsx"); self.typical = cat.load_typical(self.root / "typical.xlsx")
        self.project = {"review_dir": str(self.root)}
        self.page = "KITCHEN: 2 x UB FRIDGE 2 DOOR. FRYER UNDER EXHAUST CANOPY. OFFICE: PC MODEL A (TEST PC). CAFE: WOK."
        self.inputs = {"result": {"run": "r1", "findings": []}, "decisions": {}, "inventory": {"2": {"text": self.page}},
                       "rooms": ["Kitchen", "Office"]}

    def add(self, name, **value):
        row = {"id": f"equipment:{len(self.inputs['result']['findings'])}", "value": {"name": name, **value}, "pages": [2]}
        self.inputs["result"]["findings"].append(row)
        return row["id"]

    def resolve(self, answers=(), persist=False, book=None, typical=None):
        with patch.object(service.catalogue, "load", return_value=book or self.book), \
                patch.object(service.catalogue, "load_typical", return_value=typical or self.typical), \
                patch.object(service, "_inputs", return_value=self.inputs):
            return service.resolve(self.project, ["Kitchen", "Office"], assess=(lambda *args: list(answers)) if answers else None,
                                   persist=persist)

    def record(self, result, item_id):
        return next(row for row in result["records"] if row["id"] == item_id)

    def test_a_data_sheet_value_wins(self):
        item = self.add("Deep fryer", room="Kitchen", quantity=1)
        self.inputs["decisions"] = {item: {"run": "r1", "status": "accepted", "reviewer": "Operator", "value": {
            "name": "Deep fryer", "room": "Kitchen", "quantity": 2, "sensible_to_room_w": 0, "latent_to_room_w": 100,
            "heat_reference": "Maker data sheet"}}}
        row = self.record(self.resolve(), item)
        self.assertEqual((row["status"], row["sensible_total_w"], row["latent_total_w"]), ("data_sheet", 0, 200))
        self.assertEqual(row["use_pattern"], "WHILE_OPEN")          # no schedule entered: the type's pattern

    def test_a_printed_model_matches_only_with_its_quote(self):
        item = self.add("Test PC", room="Office", quantity=1)
        answer = {"id": item, "catalogue_id": "EQ-1", "checks": {"model": check("PC MODEL A")}}
        row = self.record(self.resolve([answer]), item)
        self.assertEqual((row["status"], row["catalogue_id"], row["sensible_w"], row["latent_w"]), ("matched_model", "EQ-1", 80, 20))
        answer["checks"]["model"] = check("PC MODEL B, invented")
        row = self.record(self.resolve([answer]), item)
        self.assertEqual(row["status"], "generic")                  # never stops: the generic placeholder rating
        self.assertIn("generic placeholder", row["source"])

    def test_one_matching_size_is_used_free_standing_unless_a_hood_is_quoted(self):
        item = self.add("Deep fryer", room="Kitchen", quantity=1)
        row = self.record(self.resolve(), item)
        self.assertEqual((row["status"], row["typical_id"], row["sensible_w"], row["latent_w"], row["heat_basis"]),
                         ("typical", "TV-3", 700, 1600, "free-standing"))
        self.assertIsNone(row["schedule"]); self.assertIsNone(row["diversity_factor"])   # the room's hours and diversity
        hooded = {"id": item, "under_hood": True, "checks": {"under_hood": check("FRYER UNDER EXHAUST CANOPY")}}
        row = self.record(self.resolve([hooded]), item)
        self.assertEqual((row["sensible_w"], row["latent_w"], row["heat_basis"]), (300, 0, "under a hood"))
        unquoted = {"id": item, "under_hood": True, "checks": {}}
        self.assertEqual(self.record(self.resolve([unquoted]), item)["sensible_w"], 700)

    def test_the_ai_chooses_between_sizes_and_continuous_items_run_every_hour(self):
        item = self.add("UB fridge", room="Kitchen", quantity=2)
        row = self.record(self.resolve(), item)
        self.assertEqual(row["status"], "generic")                   # two sizes fit, none chosen
        self.assertIn("More than one size", " ".join(row["reasons"]))
        row = self.record(self.resolve([{"id": item, "typical_id": "TV-2"}]), item)
        self.assertEqual((row["status"], row["typical_id"], row["sensible_total_w"]), ("typical", "TV-2", 1000))
        self.assertEqual((row["schedule"]["weekday"], row["diversity_factor"]), ([1.0] * 24, 1.0))

    def test_generic_value_then_no_value(self):
        item = self.add("Wok", room="Kitchen", quantity=1)
        row = self.record(self.resolve(), item)
        self.assertEqual((row["status"], row["typical_id"], row["sensible_w"]), ("generic", "TV-9", 1000))
        no_generic = {**self.typical, "records": [r for r in self.typical["records"] if r["id"] != "TV-9"]}
        result = self.resolve(typical=no_generic)
        row = self.record(result, item)
        self.assertEqual((row["status"], row["note"]), ("no_value", service.NO_VALUE_NOTE))
        self.assertFalse(result["coverage_complete"])

    def test_room_and_quantity_fall_back_and_say_so(self):
        located = self.add("Deep fryer", location="KITCHEN AREA")
        unplaced = self.add("Deep fryer")
        result = self.resolve()
        row = self.record(result, located)
        self.assertEqual((row["room"], row["quantity"]), ("Kitchen", 1))
        self.assertIn("one assumed", " ".join(row["reasons"]))
        self.assertEqual(self.record(result, unplaced)["status"], "no_room")
        answer = {"id": unplaced, "room": "Kitchen", "quantity": 2, "checks": {"quantity": check("2 x UB FRIDGE")}}
        row = self.record(self.resolve([answer]), unplaced)
        self.assertEqual((row["status"], row["room"], row["quantity"]), ("typical", "Kitchen", 2))
        self.assertIn("without a quote", " ".join(row["reasons"]))

    def test_hoods_cold_rooms_lighting_and_rejected_items_are_not_equipment_heat(self):
        hood, cold, light = self.add("Range hood", room="Kitchen"), self.add("Coolroom", room="Kitchen"), self.add("Downlight", category="Lighting")
        rejected = self.add("Deep fryer", room="Kitchen")
        self.inputs["decisions"] = {rejected: {"status": "rejected", "run": "r1"}}
        result = self.resolve()
        self.assertEqual({self.record(result, i)["status"] for i in (hood, cold, light, rejected)}, {"excluded"})

    def test_a_missing_workbook_or_failed_ai_never_stops_the_values(self):
        item = self.add("Deep fryer", room="Kitchen", quantity=1)
        broken = {"hash": "unavailable", "records": [], "diagnostics": [], "error": "Equipment catalogue cannot be read: No such file."}
        result = self.resolve(book=broken)
        self.assertEqual(self.record(result, item)["status"], "typical")
        self.assertIn("cannot be read", result["catalogue_error"])
        with patch.object(service.catalogue, "load", return_value=self.book), \
                patch.object(service.catalogue, "load_typical", return_value=self.typical), \
                patch.object(service, "_inputs", return_value=self.inputs):
            failed = service.resolve(self.project, ["Kitchen"], assess=lambda *args: (_ for _ in ()).throw(RuntimeError("timeout")))
        self.assertEqual(self.record(failed, item)["status"], "typical")
        self.assertIn("timeout", failed["assessment_error"])

    def test_saved_answers_are_reused_until_a_workbook_or_the_drawings_change(self):
        item = self.add("UB fridge", room="Kitchen", quantity=2)
        first = self.resolve([{"id": item, "typical_id": "TV-2"}], persist=True)
        again = self.resolve()
        self.assertEqual(self.record(again, item)["typical_id"], "TV-2")
        changed = {**self.typical, "hash": "edited"}
        self.assertNotEqual(first["source_hash"], self.resolve(typical=changed)["source_hash"])
        self.assertEqual(self.record(self.resolve(typical=changed), item)["status"], "generic")

    def test_incomplete_reading_is_said(self):
        self.inputs["reading_job"] = {"status": "running"}
        result = self.resolve()
        self.assertFalse(result["coverage_complete"])
        self.assertIn("isn't complete", result["reading_note"])

    def test_the_ai_check_runs_once_on_consented_pages_and_not_without_consent(self):
        from backend import ai_provider, skill_workflow_service
        item = self.add("UB fridge", room="Kitchen", quantity=2)
        (self.root / "vision_extraction_settings.json").write_text('{"owner_opt_in":true}')
        prompts = []

        class Provider:
            def propose(self, prompt, image_paths=()):
                prompts.append(prompt)
                return {"assessments": [{"id": item, "typical_id": "TV-2"}]}, {}
        patches = (patch.object(service.catalogue, "load", return_value=self.book),
                   patch.object(service.catalogue, "load_typical", return_value=self.typical),
                   patch.object(service, "_inputs", return_value=self.inputs),
                   patch.object(skill_workflow_service, "_project_paths", return_value={"root": self.root}))
        with patches[0], patches[1], patches[2], patches[3], \
                patch.object(skill_workflow_service, "_consented_page_ids", return_value={2}), \
                patch.object(ai_provider, "get", return_value=Provider()):
            first = service.assess_for_calculation(self.project, ["Kitchen"])
            second = service.assess_for_calculation(self.project, ["Kitchen"])
        self.assertEqual(self.record(first, item)["typical_id"], "TV-2")
        self.assertEqual(second["summary"], first["summary"])
        self.assertEqual(len(prompts), 1)
        self.assertIn("Never guess a model", prompts[0])
        self.assertIn("2 x UB FRIDGE", prompts[0])                      # the cited page's reading is sent
        (self.root / service.ARTIFACT).unlink()
        with patches[0], patches[1], patches[2], patches[3], \
                patch.object(skill_workflow_service, "_consented_page_ids", return_value=set()), \
                patch.object(ai_provider, "get") as provider:
            result = service.assess_for_calculation(self.project, ["Kitchen"])
        provider.assert_not_called()
        self.assertIn("consent", result["assessment_error"])

    def test_a_valued_item_reaches_a_saved_phase2_result(self):
        from types import SimpleNamespace
        from backend import model_input_resolution_service
        item = self.add("Deep fryer", room="Kitchen", quantity=2)
        project = {**self.project, "id": "catalogue-e2e"}
        artifacts = {
            "page_extraction.json": {"equipment_appliances": self.inputs["result"]},
            "page_inventory.json": {"pages": self.inputs["inventory"]},
            "building_evidence.json": {"spaces": [{"name": "Kitchen", "level_name": "Ground", "area_m2": 20,
                                                      "evidence": [{"page": 2, "excerpt": "KITCHEN 20 m2"}]}]},
            "ai_preliminary_run.json": {"manual_placeholder_proposal": {"rooms": [
                {"kind": "room", "label": "Kitchen", "level_name": "Ground", "area_m2": 20,
                 "preliminary_profile_id": "hospitality", "page": 2}]}},
        }
        for name, value in artifacts.items():
            (self.root / name).write_text(json.dumps(value))
        web = SimpleNamespace(safe_link=lambda path: str(path), update_project=lambda project: None)
        with patch.object(service.catalogue, "load", return_value=self.book), \
                patch.object(service.catalogue, "load_typical", return_value=self.typical):
            model_input_resolution_service.post(web, project, {"action": "resolve"})
            assembled = ai_preliminary_service._assemble(web, project)
            paths = ai_preliminary_service._paths(project)
            scope = ai_preliminary_service._room_scope_state(paths, assembled)
            ai_preliminary_service._confirm_room_scope(paths, {"reviewer": "Automatic room scope",
                "candidate_fingerprint": scope["candidate_fingerprint"],
                "rows": [{"key": row["key"], "include": True} for row in scope["candidates"]]})
            report = ai_preliminary_service._calculate(web, project)
        self.assertEqual(report["equipment_coverage"]["summary"]["typical"], 1)
        self.assertIn("1 typical and 0 generic equipment values", report["label"])
        sources = assembled["material"]["hourly_load_model"]["rooms"][0]["heat_sources"]
        fryer = next(source for source in sources if source["name"] == "Deep fryer")
        self.assertEqual((fryer["quantity"], fryer["watts"], fryer["latent_w"]), (2, 700, 1600))
        self.assertIn("Typical value TV-3", fryer["source"])


class CalculationTests(unittest.TestCase):
    def test_catalogue_items_keep_their_schedule_and_the_room_keeps_its_allowance_floor(self):
        building = {"spaces": [{"name": "Cafe", "level_name": "Ground", "area_m2": 100,
                                 "evidence": [{"page": 2, "excerpt": "CAFE"}]}]}
        item = {"name": "Test PC", "quantity": 2, "rated_input_w": 80, "latent_w": 20,
                "heat_to_space_factor": 1, "diversity_factor": 1, "equipment_id": "equipment:x",
                "equipment_schedule": schedule(), "evidence": [{"page": 2, "excerpt": "Model A desktop"}]}
        proposal = {"rooms": [{"kind": "room", "label": "Cafe", "level_name": "Ground", "area_m2": 100,
                               "page": 2, "preliminary_profile_id": "hospitality", "equipment_catalogue_managed": True,
                               "equipment": [item]}]}
        use = room_use_resolution.resolve(building, proposal=proposal, source_fingerprints={})
        gains = internal_gains_resolution.resolve(building, proposal=proposal, room_use=use, pack=ai_preliminary.load_pack())
        result = ai_preliminary.assemble(building, preliminary_proposal=proposal, room_use_resolution=use, internal_gains_resolution=gains)
        room = result["material"]["hourly_load_model"]["rooms"][0]
        # 160 W of listed items is below a 100 m² cafe's typical allowance: the room is topped up to it, so items with
        # no value or no room never leave it short (user decision 2026-10-20: never stop).
        self.assertEqual(len(room["heat_sources"]), 2)
        self.assertTrue(any(r.get("field") == "equipment_allowance_top_up_w" for r in result["materialized_fields"]))
        source = room["heat_sources"][0]
        contribution = equipment_load([source])
        self.assertAlmostEqual(contribution["sensible_kw"] * .5, .08)
        self.assertAlmostEqual(contribution["latent_kw"] * .5, .02)
        self.assertTrue(room["schedule_assignments"]["equipment"][source["source_id"]].startswith("equipment-"))
        first_envelope = ai_preliminary_service._phase1_input_fingerprint(result)
        changed = deepcopy(result)
        changed["material"]["hourly_load_model"]["rooms"][0]["heat_sources"][0]["watts"] = 800
        changed["material"]["schedule_library"]["schedules"].append({"schedule_id": "new-equipment-hours"})
        changed["proposal"]["equipment_catalogue_coverage"] = {"catalogue_hash": "changed"}
        self.assertEqual(first_envelope, ai_preliminary_service._phase1_input_fingerprint(changed))
        changed["material"]["hourly_load_model"]["rooms"][0]["area_m2"] = 200
        self.assertNotEqual(first_envelope, ai_preliminary_service._phase1_input_fingerprint(changed))

    def test_latent_only_equipment_still_needs_its_schedule(self):
        from ai.hourly_loads import heat_source_is_timed
        self.assertTrue(heat_source_is_timed({"quantity": 1, "watts": 0, "latent_w": 200,
                                              "space_gain_factor": 1, "diversity_factor": 1}))


class FreshnessTests(unittest.TestCase):
    def test_catalogue_change_stales_phase2_without_staling_phase1_or_rewriting_reports(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            paths = ai_preliminary_service._paths({"review_dir": folder})
            paths["current_set"].write_text(json.dumps({"input_fingerprint": "same"}))
            paths["base_report"].write_text(json.dumps({"preliminary_input_set_fingerprint": "same"}))
            paths["report"].write_text(json.dumps({"equipment_source_hash": "old", "label": "Saved result"}))
            previous = paths["report"].read_bytes()
            class Web:
                def safe_link(self, path):
                    return str(path)
            with patch.object(ai_preliminary_service, "_stale_reasons", return_value=[]), \
                 patch.object(ai_preliminary_service, "_proposal_for_resolution", return_value={"rooms": []}), \
                 patch.object(ai_preliminary_service, "_room_scope_state", return_value={}), \
                 patch.object(service, "resolve", return_value={"source_hash": "new"}):
                result = ai_preliminary_service._response(Web(), {"review_dir": folder, "id": "test"}, "workspace", True)
            self.assertTrue(result["hourly_ai_preliminary_load_report_stale"])
            self.assertFalse(result["phase1_base_room_load_report_stale"])
            self.assertEqual(paths["report"].read_bytes(), previous)


if __name__ == "__main__":
    unittest.main()
